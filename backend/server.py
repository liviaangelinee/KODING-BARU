from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends
from fastapi.responses import StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import logging
import uuid
import io
import calendar
from datetime import datetime, timezone, timedelta, date
from typing import List, Optional, Literal
from pydantic import BaseModel, Field
import bcrypt
import jwt
from bson import ObjectId

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

JWT_ALGORITHM = "HS256"
PABRIK_LIST = ["BLUTO", "PANDAAN", "IT"]

# ----------------------- Auth helpers -----------------------

def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")

def verify_password(plain: str, hashed: str) -> bool:
    return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))

def get_jwt_secret() -> str:
    return os.environ["JWT_SECRET"]

def create_access_token(user_id: str, email: str) -> str:
    payload = {"sub": user_id, "email": email,
               "exp": datetime.now(timezone.utc) + timedelta(hours=12), "type": "access"}
    return jwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)

def create_refresh_token(user_id: str) -> str:
    payload = {"sub": user_id, "exp": datetime.now(timezone.utc) + timedelta(days=7), "type": "refresh"}
    return jwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)

def set_auth_cookies(response: Response, access: str, refresh: str):
    response.set_cookie(key="access_token", value=access, httponly=True, secure=True,
                        samesite="none", max_age=43200, path="/")
    response.set_cookie(key="refresh_token", value=refresh, httponly=True, secure=True,
                        samesite="none", max_age=604800, path="/")

async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "access":
            raise HTTPException(status_code=401, detail="Token tidak valid")
        user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
        if not user:
            raise HTTPException(status_code=401, detail="User tidak ditemukan")
        user["_id"] = str(user["_id"])
        user.pop("password_hash", None)
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token kedaluwarsa")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")

# ----------------------- Models -----------------------

class LoginInput(BaseModel):
    email: str
    password: str

class UpdateNameInput(BaseModel):
    name: str

class UpdateEmailInput(BaseModel):
    new_email: str
    current_password: str

class UpdatePasswordInput(BaseModel):
    current_password: str
    new_password: str

class Variant(BaseModel):
    label: str
    harga_pabrik: float = 0   # harga beli dari pabrik (dasar perhitungan HPP/laba)
    harga_so: float = 0       # satu-satunya harga jual yang dipakai

class ProductInput(BaseModel):
    nama: str
    variants: List[Variant] = []

class CustomerInput(BaseModel):
    nama: str
    telepon: str = ""
    alamat: str = ""

class TransactionItem(BaseModel):
    product_id: str
    product_nama: str
    variant_label: str
    harga_satuan: float
    qty: float
    subtotal: float
    # Snapshot harga pabrik saat transaksi dibuat. Disimpan per item supaya laba
    # historis tetap akurat walaupun harga pabrik berubah di kemudian hari.
    harga_pabrik: float = 0
    hpp_subtotal: float = 0

class TransactionInput(BaseModel):
    tanggal: str  # YYYY-MM-DD
    customer_id: str = ""
    customer_nama: str = ""
    items: List[TransactionItem]
    # lunas = sudah dibayar | kredit = belum dibayar (piutang)
    # free  = barang gratis / promo / pemakaian sendiri -> omset Rp 0 tapi HPP tetap dihitung
    status: Literal["lunas", "kredit", "free"] = "lunas"
    # Alasan barang free (promo / sampel / pemakaian sendiri / teks manual).
    # Hanya dipakai bila status = free.
    free_alasan: str = ""
    catatan: str = ""

# Alasan bawaan untuk barang free. User tetap bisa menulis alasan sendiri (manual)
# dan alasan itu akan otomatis muncul sebagai pilihan berikutnya.
FREE_REASONS_DEFAULT = ["Promo", "Sampel", "Pemakaian Sendiri", "Lainnya"]

# ----------------------- PO Pabrik & Pengeluaran -----------------------

EXPENSE_CATEGORIES = ["BBM / Transport / Armada", "Lain-lain"]

class POItem(BaseModel):
    product_id: str = ""
    product_nama: str
    variant_label: str
    harga_pabrik: float = 0
    qty: float = 0
    subtotal: float = 0

class PurchaseOrderInput(BaseModel):
    tanggal: str  # YYYY-MM-DD
    pabrik: str
    no_po: str = ""
    items: List[POItem]
    status: Literal["lunas", "belum_lunas"] = "belum_lunas"
    catatan: str = ""

class ExpenseInput(BaseModel):
    tanggal: str  # YYYY-MM-DD
    kategori: str
    deskripsi: str = ""
    jumlah: float

class DepositInput(BaseModel):
    """Setoran uang tunai ke pemilik."""
    tanggal: str  # YYYY-MM-DD
    jumlah: float
    catatan: str = ""

# ----------------------- Auth routes -----------------------

MAX_ATTEMPTS = 5
LOCKOUT_MINUTES = 15

def get_client_ip(request: Request) -> str:
    """Real client IP. Behind the ingress/proxy, request.client.host is the proxy
    pod IP and it rotates between pods, which would split the failed-login counter
    and silently defeat the lockout. Prefer the forwarded headers."""
    xff = request.headers.get("x-forwarded-for", "")
    if xff:
        first = xff.split(",")[0].strip()
        if first:
            return first
    real = request.headers.get("x-real-ip", "").strip()
    if real:
        return real
    return request.client.host if request.client else "unknown"

@api_router.post("/auth/login")
async def login(data: LoginInput, request: Request, response: Response):
    email = data.email.lower().strip()
    ip = get_client_ip(request)
    identifier = f"{ip}:{email}"

    now = datetime.now(timezone.utc)
    rec = await db.login_attempts.find_one({"identifier": identifier})
    if rec and rec.get("count", 0) >= MAX_ATTEMPTS:
        last = rec.get("last")
        if isinstance(last, str):
            last = datetime.fromisoformat(last)
        if last and now - last < timedelta(minutes=LOCKOUT_MINUTES):
            raise HTTPException(status_code=429,
                detail=f"Terlalu banyak percobaan. Coba lagi dalam {LOCKOUT_MINUTES} menit.")

    user = await db.users.find_one({"email": email})
    if not user or not verify_password(data.password, user["password_hash"]):
        await db.login_attempts.update_one(
            {"identifier": identifier},
            {"$inc": {"count": 1}, "$set": {"last": now.isoformat()}},
            upsert=True)
        raise HTTPException(status_code=401, detail="Email atau password salah")

    await db.login_attempts.delete_one({"identifier": identifier})
    uid = str(user["_id"])
    access = create_access_token(uid, email)
    refresh = create_refresh_token(uid)
    set_auth_cookies(response, access, refresh)
    return {"id": uid, "email": email, "name": user.get("name", "Admin"), "role": user.get("role", "admin")}

@api_router.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    response.delete_cookie("refresh_token", path="/")
    return {"message": "Logout berhasil"}

@api_router.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return {"id": user["_id"], "email": user["email"], "name": user.get("name", "Admin"), "role": user.get("role", "admin")}

@api_router.put("/auth/profile/name")
async def update_name(data: UpdateNameInput, user: dict = Depends(get_current_user)):
    name = data.name.strip()
    if not name:
        raise HTTPException(status_code=422, detail="Nama tidak boleh kosong")
    await db.users.update_one({"_id": ObjectId(user["_id"])}, {"$set": {"name": name}})
    return {"id": user["_id"], "email": user["email"], "name": name, "role": user.get("role", "admin")}

@api_router.put("/auth/profile/email")
async def update_email(data: UpdateEmailInput, user: dict = Depends(get_current_user)):
    full = await db.users.find_one({"_id": ObjectId(user["_id"])})
    if not verify_password(data.current_password, full["password_hash"]):
        raise HTTPException(status_code=400, detail="Password saat ini salah")
    new_email = data.new_email.lower().strip()
    if not new_email or "@" not in new_email:
        raise HTTPException(status_code=422, detail="Email tidak valid")
    dup = await db.users.find_one({"email": new_email, "_id": {"$ne": ObjectId(user["_id"])}})
    if dup:
        raise HTTPException(status_code=400, detail="Email sudah digunakan")
    await db.users.update_one({"_id": ObjectId(user["_id"])}, {"$set": {"email": new_email}})
    return {"id": user["_id"], "email": new_email, "name": full.get("name", "Admin"), "role": full.get("role", "admin")}

@api_router.put("/auth/profile/password")
async def update_password_ep(data: UpdatePasswordInput, user: dict = Depends(get_current_user)):
    full = await db.users.find_one({"_id": ObjectId(user["_id"])})
    if not verify_password(data.current_password, full["password_hash"]):
        raise HTTPException(status_code=400, detail="Password saat ini salah")
    if len(data.new_password) < 6:
        raise HTTPException(status_code=422, detail="Password baru minimal 6 karakter")
    await db.users.update_one({"_id": ObjectId(user["_id"])}, {"$set": {"password_hash": hash_password(data.new_password)}})
    return {"message": "Password berhasil diperbarui"}

@api_router.post("/auth/refresh")
async def refresh_token(request: Request, response: Response):
    token = request.cookies.get("refresh_token")
    if not token:
        raise HTTPException(status_code=401, detail="Tidak ada refresh token")
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "refresh":
            raise HTTPException(status_code=401, detail="Token tidak valid")
        user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
        if not user:
            raise HTTPException(status_code=401, detail="User tidak ditemukan")
        access = create_access_token(str(user["_id"]), user["email"])
        response.set_cookie(key="access_token", value=access, httponly=True, secure=True,
                            samesite="none", max_age=43200, path="/")
        return {"message": "ok"}
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")

# ----------------------- Products -----------------------

@api_router.get("/products")
async def list_products(user: dict = Depends(get_current_user)):
    items = await db.products.find({}, {"_id": 0}).sort("nama", 1).to_list(1000)
    return items

@api_router.post("/products")
async def create_product(data: ProductInput, user: dict = Depends(get_current_user)):
    doc = {"id": str(uuid.uuid4()), "nama": data.nama,
           "variants": [v.model_dump() for v in data.variants],
           "created_at": datetime.now(timezone.utc).isoformat()}
    await db.products.insert_one(doc.copy())
    doc.pop("_id", None)
    return doc

@api_router.put("/products/{product_id}")
async def update_product(product_id: str, data: ProductInput, user: dict = Depends(get_current_user)):
    res = await db.products.update_one({"id": product_id}, {"$set": {
        "nama": data.nama, "variants": [v.model_dump() for v in data.variants]}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    return await db.products.find_one({"id": product_id}, {"_id": 0})

@api_router.delete("/products/{product_id}")
async def delete_product(product_id: str, user: dict = Depends(get_current_user)):
    res = await db.products.delete_one({"id": product_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Produk tidak ditemukan")
    return {"message": "Produk dihapus"}

# ----------------------- Customers -----------------------

@api_router.get("/customers")
async def list_customers(user: dict = Depends(get_current_user)):
    items = await db.customers.find({}, {"_id": 0}).sort("nama", 1).to_list(1000)
    return items

@api_router.post("/customers")
async def create_customer(data: CustomerInput, user: dict = Depends(get_current_user)):
    doc = {"id": str(uuid.uuid4()), **data.model_dump(),
           "created_at": datetime.now(timezone.utc).isoformat()}
    await db.customers.insert_one(doc.copy())
    doc.pop("_id", None)
    return doc

@api_router.put("/customers/{customer_id}")
async def update_customer(customer_id: str, data: CustomerInput, user: dict = Depends(get_current_user)):
    res = await db.customers.update_one({"id": customer_id}, {"$set": data.model_dump()})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Customer tidak ditemukan")
    return await db.customers.find_one({"id": customer_id}, {"_id": 0})

@api_router.delete("/customers/{customer_id}")
async def delete_customer(customer_id: str, user: dict = Depends(get_current_user)):
    res = await db.customers.delete_one({"id": customer_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Customer tidak ditemukan")
    return {"message": "Customer dihapus"}

# ----------------------- Transactions -----------------------

def build_txn_query(date_from, date_to, customer_id, status):
    q = {}
    if date_from and date_to:
        q["tanggal"] = {"$gte": date_from, "$lte": date_to}
    elif date_from:
        q["tanggal"] = {"$gte": date_from}
    elif date_to:
        q["tanggal"] = {"$lte": date_to}
    if customer_id:
        q["customer_id"] = customer_id
    if status:
        q["status"] = status
    return q

async def enrich_items_with_hpp(items: List[TransactionItem]) -> List[dict]:
    """Isi harga_pabrik + hpp_subtotal pada setiap item transaksi.

    Harga pabrik diambil dari varian produk saat transaksi dibuat, lalu di-SNAPSHOT
    ke dalam item. Dengan begitu laba periode lampau tidak berubah ketika harga
    pabrik dinaikkan di menu Produk.
    """
    pids = list({i.product_id for i in items if i.product_id})
    prods = await db.products.find({"id": {"$in": pids}}, {"_id": 0}).to_list(1000)
    pmap = {p["id"]: p for p in prods}
    out = []
    for i in items:
        d = i.model_dump()
        hp = float(d.get("harga_pabrik") or 0)
        if hp <= 0:
            p = pmap.get(i.product_id)
            if p:
                v = next((x for x in p.get("variants", []) if x.get("label") == i.variant_label), None)
                if v:
                    hp = float(v.get("harga_pabrik") or 0)
        d["harga_pabrik"] = hp
        d["hpp_subtotal"] = hp * float(d.get("qty") or 0)
        out.append(d)
    return out

def txn_hpp(t: dict) -> float:
    """HPP sebuah transaksi, tahan terhadap dokumen lama tanpa field total_hpp."""
    if t.get("total_hpp") is not None:
        return float(t["total_hpp"])
    return sum(float(it.get("hpp_subtotal") or 0) for it in t.get("items", []))

@api_router.get("/transactions")
async def list_transactions(date_from: Optional[str] = None, date_to: Optional[str] = None,
                            customer_id: Optional[str] = None, status: Optional[str] = None,
                            user: dict = Depends(get_current_user)):
    q = build_txn_query(date_from, date_to, customer_id, status)
    items = await db.transactions.find(q, {"_id": 0}).sort("tanggal", -1).to_list(5000)
    for t in items:
        hpp = txn_hpp(t)
        t["total_hpp"] = hpp
        t["laba_kotor"] = float(t.get("total") or 0) - hpp
    return items

def normalize_txn(data: TransactionInput):
    """Aturan khusus status 'free'.

    Barang free (promo, sampel, atau pemakaian sendiri) tidak menghasilkan omset,
    jadi harga jualnya dipaksa Rp 0. Tetapi HPP-nya tetap dihitung di
    enrich_items_with_hpp, sehingga modal barang itu tetap mengurangi laba.
    Customer boleh kosong untuk pemakaian sendiri.
    """
    items = data.items
    alasan = (data.free_alasan or "").strip()
    if data.status == "free":
        for i in items:
            i.harga_satuan = 0
            i.subtotal = 0
        if not alasan:
            alasan = "Pemakaian Sendiri"
    else:
        alasan = ""
    nama = (data.customer_nama or "").strip()
    if not nama:
        if data.status == "free":
            nama = "Pemakaian Sendiri"
        else:
            raise HTTPException(status_code=422, detail="Customer wajib dipilih")
    return items, nama, alasan

@api_router.get("/transactions/free-reasons")
async def free_reasons(user: dict = Depends(get_current_user)):
    """Daftar alasan barang free: bawaan + alasan manual yang pernah dipakai."""
    used = await db.transactions.distinct("free_alasan", {"status": "free"})
    extra = sorted({(u or "").strip() for u in used if (u or "").strip()} - set(FREE_REASONS_DEFAULT))
    return FREE_REASONS_DEFAULT + extra

@api_router.post("/transactions")
async def create_transaction(data: TransactionInput, user: dict = Depends(get_current_user)):
    items, customer_nama, free_alasan = normalize_txn(data)
    total = sum(i.subtotal for i in items)
    enriched = await enrich_items_with_hpp(items)
    total_hpp = sum(i["hpp_subtotal"] for i in enriched)
    doc = {"id": str(uuid.uuid4()), "tanggal": data.tanggal, "customer_id": data.customer_id,
           "customer_nama": customer_nama, "items": enriched,
           "total": total, "total_hpp": total_hpp, "laba_kotor": total - total_hpp,
           "status": data.status, "free_alasan": free_alasan, "catatan": data.catatan,
           "created_at": datetime.now(timezone.utc).isoformat()}
    await db.transactions.insert_one(doc.copy())
    doc.pop("_id", None)
    return doc

@api_router.put("/transactions/{txn_id}")
async def update_transaction(txn_id: str, data: TransactionInput, user: dict = Depends(get_current_user)):
    items, customer_nama, free_alasan = normalize_txn(data)
    total = sum(i.subtotal for i in items)
    enriched = await enrich_items_with_hpp(items)
    total_hpp = sum(i["hpp_subtotal"] for i in enriched)
    res = await db.transactions.update_one({"id": txn_id}, {"$set": {
        "tanggal": data.tanggal, "customer_id": data.customer_id, "customer_nama": customer_nama,
        "items": enriched, "total": total, "total_hpp": total_hpp,
        "laba_kotor": total - total_hpp,
        "status": data.status, "free_alasan": free_alasan, "catatan": data.catatan}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Transaksi tidak ditemukan")
    return await db.transactions.find_one({"id": txn_id}, {"_id": 0})

@api_router.patch("/transactions/{txn_id}/status")
async def update_status(txn_id: str, body: dict, user: dict = Depends(get_current_user)):
    status = body.get("status")
    if status not in ("lunas", "kredit", "free"):
        raise HTTPException(status_code=400, detail="Status tidak valid")
    txn = await db.transactions.find_one({"id": txn_id})
    if not txn:
        raise HTTPException(status_code=404, detail="Transaksi tidak ditemukan")

    if txn.get("status") == "free" and status != "free":
        # Harga jual barang free tidak pernah dicatat (selalu 0), jadi tidak bisa
        # dipulihkan. Lebih aman minta dibuat ulang daripada menyimpan angka salah.
        raise HTTPException(status_code=400,
            detail="Transaksi Free tidak bisa diubah jadi Lunas/Kredit. Hapus lalu buat transaksi baru.")

    if status == "free":
        items = txn.get("items", [])
        for it in items:
            it["harga_satuan"] = 0
            it["subtotal"] = 0
        total_hpp = sum(float(i.get("hpp_subtotal") or 0) for i in items)
        alasan = (body.get("free_alasan") or txn.get("free_alasan") or "").strip() or "Pemakaian Sendiri"
        await db.transactions.update_one({"id": txn_id}, {"$set": {
            "status": "free", "items": items, "total": 0, "free_alasan": alasan,
            "total_hpp": total_hpp, "laba_kotor": -total_hpp}})
    else:
        await db.transactions.update_one({"id": txn_id}, {"$set": {"status": status, "free_alasan": ""}})
    return {"message": "Status diperbarui"}

@api_router.delete("/transactions/{txn_id}")
async def delete_transaction(txn_id: str, user: dict = Depends(get_current_user)):
    res = await db.transactions.delete_one({"id": txn_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Transaksi tidak ditemukan")
    return {"message": "Transaksi dihapus"}

# ----------------------- Dashboard -----------------------

@api_router.get("/dashboard/stats")
async def dashboard_stats(user: dict = Depends(get_current_user)):
    txns = await db.transactions.find({}, {"_id": 0}).to_list(10000)
    total_omset = sum(t["total"] for t in txns)
    total_txn = len(txns)
    sisa_tagihan = sum(t["total"] for t in txns if t["status"] == "kredit")
    total_customer = await db.customers.count_documents({})

    # omset trend by date (last 30 days that exist)
    by_date = {}
    for t in txns:
        by_date[t["tanggal"]] = by_date.get(t["tanggal"], 0) + t["total"]
    trend = [{"tanggal": k, "omset": v} for k, v in sorted(by_date.items())][-30:]

    # top products by omset
    prod = {}
    prod_qty = {}
    for t in txns:
        for it in t["items"]:
            key = f'{it["product_nama"]} {it["variant_label"]}'
            prod[key] = prod.get(key, 0) + it["subtotal"]
            prod_qty[key] = prod_qty.get(key, 0) + it["qty"]
    top_products = sorted(
        [{"nama": k, "omset": v, "qty": prod_qty[k]} for k, v in prod.items()],
        key=lambda x: x["omset"], reverse=True)[:8]

    # outstanding per customer
    outstanding = {}
    for t in txns:
        if t["status"] == "kredit":
            outstanding[t["customer_nama"]] = outstanding.get(t["customer_nama"], 0) + t["total"]
    outstanding_list = sorted(
        [{"customer": k, "sisa": v} for k, v in outstanding.items()],
        key=lambda x: x["sisa"], reverse=True)

    # Ringkasan keuangan bulan berjalan
    now = datetime.now(timezone.utc)
    prefix = f"{now.year:04d}-{now.month:02d}"
    laba = await compute_laba(prefix)

    # Setoran ke pemilik (hari ini & bulan berjalan)
    hari_ini = now.strftime("%Y-%m-%d")
    deps_bulan = await db.deposits.find({"tanggal": {"$regex": f"^{prefix}"}}, {"_id": 0}).to_list(5000)
    setoran_bulan = sum(float(d.get("jumlah") or 0) for d in deps_bulan)
    setoran_hari_ini = sum(float(d.get("jumlah") or 0) for d in deps_bulan if d.get("tanggal") == hari_ini)

    # Modal barang free (promo / pemakaian sendiri) bulan berjalan
    free_hpp_bulan = sum(txn_hpp(t) for t in txns
                         if t.get("status") == "free" and str(t.get("tanggal", "")).startswith(prefix))

    return {"total_omset": total_omset, "total_txn": total_txn, "sisa_tagihan": sisa_tagihan,
            "setoran_hari_ini": setoran_hari_ini, "setoran_bulan": setoran_bulan,
            "free_hpp_bulan": free_hpp_bulan, "tanggal_hari_ini": hari_ini,
            "total_customer": total_customer, "trend": trend, "top_products": top_products,
            "outstanding": outstanding_list,
            "bulan_ini": prefix,
            "omset_bulan": laba["omset"],
            "hpp_bulan": laba["hpp"],
            "laba_kotor_bulan": laba["laba_kotor"],
            "pengeluaran_bulan": laba["pengeluaran"],
            "laba_bersih_bulan": laba["laba_bersih"],
            "po_pabrik_bulan": laba["total_po_pabrik"],
            "hutang_pabrik": laba["hutang_pabrik"]}

@api_router.get("/rekap/monthly")
async def rekap_monthly(year: int, month: int, user: dict = Depends(get_current_user)):
    if month < 1 or month > 12:
        raise HTTPException(status_code=422, detail="Bulan harus antara 1 dan 12")
    if year < 2000 or year > 2100:
        raise HTTPException(status_code=422, detail="Tahun tidak valid")
    prefix = f"{year:04d}-{month:02d}"
    txns = await db.transactions.find({"tanggal": {"$regex": f"^{prefix}"}}, {"_id": 0}).to_list(10000)
    total_omset = sum(t["total"] for t in txns)
    total_txn = len(txns)
    sisa = sum(t["total"] for t in txns if t["status"] == "kredit")
    lunas = total_omset - sisa

    prod = {}
    for t in txns:
        for it in t["items"]:
            key = f'{it["product_nama"]} {it["variant_label"]}'
            if key not in prod:
                prod[key] = {"nama": key, "qty": 0, "omset": 0}
            prod[key]["qty"] += it["qty"]
            prod[key]["omset"] += it["subtotal"]
    per_product = sorted(prod.values(), key=lambda x: x["omset"], reverse=True)

    cust = {}
    for t in txns:
        if t["customer_nama"] not in cust:
            cust[t["customer_nama"]] = {"customer": t["customer_nama"], "omset": 0, "sisa": 0, "txn": 0}
        cust[t["customer_nama"]]["omset"] += t["total"]
        cust[t["customer_nama"]]["txn"] += 1
        if t["status"] == "kredit":
            cust[t["customer_nama"]]["sisa"] += t["total"]
    per_customer = sorted(cust.values(), key=lambda x: x["omset"], reverse=True)

    daily = {}
    for t in txns:
        daily[t["tanggal"]] = daily.get(t["tanggal"], 0) + t["total"]
    per_day = [{"tanggal": k, "omset": v} for k, v in sorted(daily.items())]

    return {"total_omset": total_omset, "total_txn": total_txn, "sisa_tagihan": sisa, "lunas": lunas,
            "per_product": per_product, "per_customer": per_customer, "per_day": per_day}

# ----------------------- PO Pabrik (Pembelian) -----------------------

def build_po_query(date_from, date_to, status, pabrik):
    q = {}
    if date_from and date_to:
        q["tanggal"] = {"$gte": date_from, "$lte": date_to}
    elif date_from:
        q["tanggal"] = {"$gte": date_from}
    elif date_to:
        q["tanggal"] = {"$lte": date_to}
    if status:
        q["status"] = status
    if pabrik:
        q["pabrik"] = pabrik
    return q

@api_router.get("/purchase-orders")
async def list_purchase_orders(date_from: Optional[str] = None, date_to: Optional[str] = None,
                               status: Optional[str] = None, pabrik: Optional[str] = None,
                               user: dict = Depends(get_current_user)):
    q = build_po_query(date_from, date_to, status, pabrik)
    return await db.purchase_orders.find(q, {"_id": 0}).sort("tanggal", -1).to_list(5000)

@api_router.get("/purchase-orders/pabrik-list")
async def pabrik_list(user: dict = Depends(get_current_user)):
    """Daftar pabrik: BLUTO, PANDAAN, IT (plus nama lain yang pernah dipakai)."""
    saved = await db.purchase_orders.distinct("pabrik")
    merged = list(PABRIK_LIST)
    for n in saved:
        if n and n not in merged:
            merged.append(n)
    return merged

@api_router.post("/purchase-orders")
async def create_purchase_order(data: PurchaseOrderInput, user: dict = Depends(get_current_user)):
    if not data.pabrik.strip():
        raise HTTPException(status_code=422, detail="Nama pabrik wajib diisi")
    if not data.items:
        raise HTTPException(status_code=422, detail="Minimal satu item pembelian")
    items = [i.model_dump() for i in data.items]
    total = sum(float(i["harga_pabrik"]) * float(i["qty"]) for i in items)
    for i in items:
        i["subtotal"] = float(i["harga_pabrik"]) * float(i["qty"])
    doc = {"id": str(uuid.uuid4()), "tanggal": data.tanggal, "pabrik": data.pabrik.strip(),
           "no_po": data.no_po.strip(), "items": items, "total": total,
           "status": data.status, "catatan": data.catatan,
           "created_at": datetime.now(timezone.utc).isoformat()}
    await db.purchase_orders.insert_one(doc.copy())
    doc.pop("_id", None)
    return doc

@api_router.put("/purchase-orders/{po_id}")
async def update_purchase_order(po_id: str, data: PurchaseOrderInput, user: dict = Depends(get_current_user)):
    items = [i.model_dump() for i in data.items]
    for i in items:
        i["subtotal"] = float(i["harga_pabrik"]) * float(i["qty"])
    total = sum(i["subtotal"] for i in items)
    res = await db.purchase_orders.update_one({"id": po_id}, {"$set": {
        "tanggal": data.tanggal, "pabrik": data.pabrik.strip(), "no_po": data.no_po.strip(),
        "items": items, "total": total, "status": data.status, "catatan": data.catatan}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="PO tidak ditemukan")
    return await db.purchase_orders.find_one({"id": po_id}, {"_id": 0})

@api_router.patch("/purchase-orders/{po_id}/status")
async def update_po_status(po_id: str, body: dict, user: dict = Depends(get_current_user)):
    status = body.get("status")
    if status not in ("lunas", "belum_lunas"):
        raise HTTPException(status_code=400, detail="Status tidak valid")
    res = await db.purchase_orders.update_one({"id": po_id}, {"$set": {"status": status}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="PO tidak ditemukan")
    return {"message": "Status pembayaran PO diperbarui"}

@api_router.delete("/purchase-orders/{po_id}")
async def delete_purchase_order(po_id: str, user: dict = Depends(get_current_user)):
    res = await db.purchase_orders.delete_one({"id": po_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="PO tidak ditemukan")
    return {"message": "PO dihapus"}

@api_router.get("/purchases/report")
async def purchases_report(year: int, month: int, user: dict = Depends(get_current_user)):
    if month < 1 or month > 12:
        raise HTTPException(status_code=422, detail="Bulan harus antara 1 dan 12")
    if year < 2000 or year > 2100:
        raise HTTPException(status_code=422, detail="Tahun tidak valid")
    prefix = f"{year:04d}-{month:02d}"
    pos = await db.purchase_orders.find({"tanggal": {"$regex": f"^{prefix}"}}, {"_id": 0}).to_list(10000)

    total_pembelian = sum(float(p.get("total") or 0) for p in pos)
    jumlah_po = len(pos)
    belum_bayar = sum(float(p.get("total") or 0) for p in pos if p.get("status") == "belum_lunas")
    sudah_bayar = total_pembelian - belum_bayar

    prod = {}
    for p in pos:
        for it in p.get("items", []):
            key = f'{it.get("product_nama","")} {it.get("variant_label","")}'.strip()
            if key not in prod:
                prod[key] = {"nama": key, "qty": 0, "total": 0}
            prod[key]["qty"] += float(it.get("qty") or 0)
            prod[key]["total"] += float(it.get("subtotal") or 0)
    per_product = sorted(prod.values(), key=lambda x: x["total"], reverse=True)

    pab = {}
    for p in pos:
        nm = p.get("pabrik", "-")
        if nm not in pab:
            pab[nm] = {"pabrik": nm, "total": 0, "po": 0, "hutang": 0}
        pab[nm]["total"] += float(p.get("total") or 0)
        pab[nm]["po"] += 1
        if p.get("status") == "belum_lunas":
            pab[nm]["hutang"] += float(p.get("total") or 0)
    per_pabrik = sorted(pab.values(), key=lambda x: x["total"], reverse=True)

    daily = {}
    for p in pos:
        daily[p["tanggal"]] = daily.get(p["tanggal"], 0) + float(p.get("total") or 0)
    per_day = [{"tanggal": k, "total": v} for k, v in sorted(daily.items())]

    # hutang total ke pabrik (semua periode, bukan hanya bulan ini)
    all_unpaid = await db.purchase_orders.find({"status": "belum_lunas"}, {"_id": 0}).to_list(10000)
    hutang_total = sum(float(p.get("total") or 0) for p in all_unpaid)

    return {"total_pembelian": total_pembelian, "jumlah_po": jumlah_po,
            "belum_bayar": belum_bayar, "sudah_bayar": sudah_bayar,
            "hutang_total": hutang_total, "per_product": per_product,
            "per_pabrik": per_pabrik, "per_day": per_day}

# ----------------------- Pengeluaran Operasional -----------------------

@api_router.get("/expenses/categories")
async def expense_categories(user: dict = Depends(get_current_user)):
    saved = await db.expenses.distinct("kategori")
    merged = list(EXPENSE_CATEGORIES)
    for s in saved:
        if s and s not in merged:
            merged.append(s)
    return merged

@api_router.get("/expenses")
async def list_expenses(date_from: Optional[str] = None, date_to: Optional[str] = None,
                        kategori: Optional[str] = None, user: dict = Depends(get_current_user)):
    q = {}
    if date_from and date_to:
        q["tanggal"] = {"$gte": date_from, "$lte": date_to}
    elif date_from:
        q["tanggal"] = {"$gte": date_from}
    elif date_to:
        q["tanggal"] = {"$lte": date_to}
    if kategori:
        q["kategori"] = kategori
    return await db.expenses.find(q, {"_id": 0}).sort("tanggal", -1).to_list(5000)

@api_router.post("/expenses")
async def create_expense(data: ExpenseInput, user: dict = Depends(get_current_user)):
    if not data.kategori.strip():
        raise HTTPException(status_code=422, detail="Kategori wajib diisi")
    if float(data.jumlah) <= 0:
        raise HTTPException(status_code=422, detail="Jumlah pengeluaran harus lebih dari 0")
    doc = {"id": str(uuid.uuid4()), "tanggal": data.tanggal, "kategori": data.kategori.strip(),
           "deskripsi": data.deskripsi.strip(), "jumlah": float(data.jumlah),
           "created_at": datetime.now(timezone.utc).isoformat()}
    await db.expenses.insert_one(doc.copy())
    doc.pop("_id", None)
    return doc

@api_router.put("/expenses/{exp_id}")
async def update_expense(exp_id: str, data: ExpenseInput, user: dict = Depends(get_current_user)):
    if float(data.jumlah) <= 0:
        raise HTTPException(status_code=422, detail="Jumlah pengeluaran harus lebih dari 0")
    res = await db.expenses.update_one({"id": exp_id}, {"$set": {
        "tanggal": data.tanggal, "kategori": data.kategori.strip(),
        "deskripsi": data.deskripsi.strip(), "jumlah": float(data.jumlah)}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Pengeluaran tidak ditemukan")
    return await db.expenses.find_one({"id": exp_id}, {"_id": 0})

@api_router.delete("/expenses/{exp_id}")
async def delete_expense(exp_id: str, user: dict = Depends(get_current_user)):
    res = await db.expenses.delete_one({"id": exp_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Pengeluaran tidak ditemukan")
    return {"message": "Pengeluaran dihapus"}

# ----------------------- Laba / Laporan Keuangan -----------------------

async def variants_without_cost():
    """Varian produk yang belum diisi Harga Pabrik -> laba kotor belum akurat."""
    prods = await db.products.find({}, {"_id": 0}).to_list(1000)
    missing = []
    for p in prods:
        for v in p.get("variants", []):
            if float(v.get("harga_pabrik") or 0) <= 0:
                missing.append({"produk": p.get("nama", ""), "varian": v.get("label", "")})
    return missing

async def compute_laba(prefix: str):
    """Hitung laba untuk periode 'YYYY-MM'.

    Laba Kotor  = Omset Penjualan - HPP (harga pabrik barang yang TERJUAL)
    Laba Bersih = Laba Kotor - Pengeluaran Operasional
    """
    txns = await db.transactions.find({"tanggal": {"$regex": f"^{prefix}"}}, {"_id": 0}).to_list(10000)
    exps = await db.expenses.find({"tanggal": {"$regex": f"^{prefix}"}}, {"_id": 0}).to_list(10000)
    pos = await db.purchase_orders.find({"tanggal": {"$regex": f"^{prefix}"}}, {"_id": 0}).to_list(10000)

    omset = sum(float(t.get("total") or 0) for t in txns)
    hpp = sum(txn_hpp(t) for t in txns)
    laba_kotor = omset - hpp
    pengeluaran = sum(float(e.get("jumlah") or 0) for e in exps)
    laba_bersih = laba_kotor - pengeluaran

    kat = {}
    for e in exps:
        k = e.get("kategori", "Lain-lain")
        kat[k] = kat.get(k, 0) + float(e.get("jumlah") or 0)
    per_kategori = sorted([{"kategori": k, "jumlah": v} for k, v in kat.items()],
                          key=lambda x: x["jumlah"], reverse=True)

    prod = {}
    for t in txns:
        for it in t.get("items", []):
            key = f'{it.get("product_nama","")} {it.get("variant_label","")}'.strip()
            if key not in prod:
                prod[key] = {"nama": key, "qty": 0, "omset": 0, "hpp": 0, "laba": 0}
            prod[key]["qty"] += float(it.get("qty") or 0)
            prod[key]["omset"] += float(it.get("subtotal") or 0)
            prod[key]["hpp"] += float(it.get("hpp_subtotal") or 0)
    for v in prod.values():
        v["laba"] = v["omset"] - v["hpp"]
    per_product = sorted(prod.values(), key=lambda x: x["laba"], reverse=True)

    daily = {}
    for t in txns:
        d = daily.setdefault(t["tanggal"], {"tanggal": t["tanggal"], "omset": 0, "hpp": 0,
                                            "pengeluaran": 0, "laba": 0})
        d["omset"] += float(t.get("total") or 0)
        d["hpp"] += txn_hpp(t)
    for e in exps:
        d = daily.setdefault(e["tanggal"], {"tanggal": e["tanggal"], "omset": 0, "hpp": 0,
                                            "pengeluaran": 0, "laba": 0})
        d["pengeluaran"] += float(e.get("jumlah") or 0)
    for d in daily.values():
        d["laba"] = d["omset"] - d["hpp"] - d["pengeluaran"]
    per_day = [daily[k] for k in sorted(daily.keys())]

    all_unpaid = await db.purchase_orders.find({"status": "belum_lunas"}, {"_id": 0}).to_list(10000)

    return {
        "omset": omset,
        "hpp": hpp,
        "laba_kotor": laba_kotor,
        "pengeluaran": pengeluaran,
        "laba_bersih": laba_bersih,
        "margin_kotor_pct": round(laba_kotor / omset * 100, 2) if omset else 0,
        "margin_bersih_pct": round(laba_bersih / omset * 100, 2) if omset else 0,
        "total_txn": len(txns),
        "total_po_pabrik": sum(float(p.get("total") or 0) for p in pos),
        "jumlah_po": len(pos),
        "hutang_pabrik": sum(float(p.get("total") or 0) for p in all_unpaid),
        "per_kategori": per_kategori,
        "per_product": per_product,
        "per_day": per_day,
        "varian_tanpa_harga_pabrik": await variants_without_cost(),
    }

@api_router.get("/laba")
async def laba_report(year: int, month: int, user: dict = Depends(get_current_user)):
    if month < 1 or month > 12:
        raise HTTPException(status_code=422, detail="Bulan harus antara 1 dan 12")
    if year < 2000 or year > 2100:
        raise HTTPException(status_code=422, detail="Tahun tidak valid")
    return await compute_laba(f"{year:04d}-{month:02d}")

BULAN_SINGKAT = ["", "Jan", "Feb", "Mar", "Apr", "Mei", "Jun",
                 "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]

async def compute_laba_tahunan(year: int):
    """Perbandingan laba 12 bulan dalam satu tahun.

    Semua data setahun diambil sekali lalu dikelompokkan per bulan, jadi hanya
    3 query ke database (bukan 36).
    """
    rx = f"^{year:04d}-"
    txns = await db.transactions.find({"tanggal": {"$regex": rx}}, {"_id": 0}).to_list(50000)
    exps = await db.expenses.find({"tanggal": {"$regex": rx}}, {"_id": 0}).to_list(50000)
    pos = await db.purchase_orders.find({"tanggal": {"$regex": rx}}, {"_id": 0}).to_list(50000)

    bulanan = {
        m: {"bulan": m, "nama_bulan": BULAN_ID[m], "nama_singkat": BULAN_SINGKAT[m],
            "omset": 0, "hpp": 0, "laba_kotor": 0, "pengeluaran": 0, "laba_bersih": 0,
            "total_txn": 0, "po_pabrik": 0, "margin_bersih_pct": 0}
        for m in range(1, 13)
    }

    def bulan_dari(tanggal):
        try:
            return int(str(tanggal)[5:7])
        except (ValueError, TypeError):
            return 0

    for t in txns:
        m = bulan_dari(t.get("tanggal"))
        if m in bulanan:
            bulanan[m]["omset"] += float(t.get("total") or 0)
            bulanan[m]["hpp"] += txn_hpp(t)
            bulanan[m]["total_txn"] += 1
    for e in exps:
        m = bulan_dari(e.get("tanggal"))
        if m in bulanan:
            bulanan[m]["pengeluaran"] += float(e.get("jumlah") or 0)
    for p in pos:
        m = bulan_dari(p.get("tanggal"))
        if m in bulanan:
            bulanan[m]["po_pabrik"] += float(p.get("total") or 0)

    for b in bulanan.values():
        b["laba_kotor"] = b["omset"] - b["hpp"]
        b["laba_bersih"] = b["laba_kotor"] - b["pengeluaran"]
        b["margin_bersih_pct"] = round(b["laba_bersih"] / b["omset"] * 100, 2) if b["omset"] else 0

    per_bulan = [bulanan[m] for m in range(1, 13)]
    aktif = [b for b in per_bulan if b["omset"] or b["pengeluaran"] or b["po_pabrik"]]

    total_omset = sum(b["omset"] for b in per_bulan)
    total_hpp = sum(b["hpp"] for b in per_bulan)
    total_pengeluaran = sum(b["pengeluaran"] for b in per_bulan)
    total_laba_kotor = total_omset - total_hpp
    total_laba_bersih = total_laba_kotor - total_pengeluaran

    terbaik = max(aktif, key=lambda x: x["laba_bersih"]) if aktif else None
    terburuk = min(aktif, key=lambda x: x["laba_bersih"]) if aktif else None

    return {
        "tahun": year,
        "per_bulan": per_bulan,
        "total_omset": total_omset,
        "total_hpp": total_hpp,
        "total_laba_kotor": total_laba_kotor,
        "total_pengeluaran": total_pengeluaran,
        "total_laba_bersih": total_laba_bersih,
        "total_po_pabrik": sum(b["po_pabrik"] for b in per_bulan),
        "total_txn": sum(b["total_txn"] for b in per_bulan),
        "margin_bersih_pct": round(total_laba_bersih / total_omset * 100, 2) if total_omset else 0,
        "bulan_aktif": len(aktif),
        "rata_laba_bersih": round(total_laba_bersih / len(aktif), 2) if aktif else 0,
        "bulan_terbaik": terbaik,
        "bulan_terburuk": terburuk,
    }

@api_router.get("/laba/yearly")
async def laba_tahunan(year: int, user: dict = Depends(get_current_user)):
    if year < 2000 or year > 2100:
        raise HTTPException(status_code=422, detail="Tahun tidak valid")
    return await compute_laba_tahunan(year)

# ----------------------- Setoran ke Pemilik -----------------------

@api_router.get("/deposits")
async def list_deposits(date_from: Optional[str] = None, date_to: Optional[str] = None,
                        user: dict = Depends(get_current_user)):
    q = {}
    if date_from and date_to:
        q["tanggal"] = {"$gte": date_from, "$lte": date_to}
    elif date_from:
        q["tanggal"] = {"$gte": date_from}
    elif date_to:
        q["tanggal"] = {"$lte": date_to}
    return await db.deposits.find(q, {"_id": 0}).sort("tanggal", -1).to_list(5000)

@api_router.post("/deposits")
async def create_deposit(data: DepositInput, user: dict = Depends(get_current_user)):
    if float(data.jumlah) <= 0:
        raise HTTPException(status_code=422, detail="Jumlah setoran harus lebih dari 0")
    doc = {"id": str(uuid.uuid4()), "tanggal": data.tanggal, "jumlah": float(data.jumlah),
           "catatan": data.catatan.strip(),
           "created_at": datetime.now(timezone.utc).isoformat()}
    await db.deposits.insert_one(doc.copy())
    doc.pop("_id", None)
    return doc

@api_router.put("/deposits/{dep_id}")
async def update_deposit(dep_id: str, data: DepositInput, user: dict = Depends(get_current_user)):
    if float(data.jumlah) <= 0:
        raise HTTPException(status_code=422, detail="Jumlah setoran harus lebih dari 0")
    res = await db.deposits.update_one({"id": dep_id}, {"$set": {
        "tanggal": data.tanggal, "jumlah": float(data.jumlah), "catatan": data.catatan.strip()}})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Setoran tidak ditemukan")
    return await db.deposits.find_one({"id": dep_id}, {"_id": 0})

@api_router.delete("/deposits/{dep_id}")
async def delete_deposit(dep_id: str, user: dict = Depends(get_current_user)):
    res = await db.deposits.delete_one({"id": dep_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Setoran tidak ditemukan")
    return {"message": "Setoran dihapus"}

# ----------------------- Laporan Harian -----------------------

def _range_query(date_from, date_to):
    return {"tanggal": {"$gte": date_from, "$lte": date_to}}

async def compute_harian(date_from: str, date_to: str):
    """Rekap per hari: penjualan, pembelian, operasional, laba bersih, setoran.

    Definisi yang dipakai:
      penjualan   = total semua transaksi (Lunas + Kredit). Free = Rp 0.
      kas_masuk   = hanya transaksi Lunas (Kredit belum jadi uang, Free tidak ada uang)
      pembelian   = total PO ke pabrik pada tanggal tersebut
      pembayaran_pabrik = PO yang statusnya sudah dibayar (dianggap dibayar pada tanggal PO,
                    karena PO tidak punya kolom tanggal bayar terpisah)
      operasional = pengeluaran operasional
      laba_bersih = (penjualan - HPP) - operasional
      seharusnya_disetor = kas_masuk - operasional - pembayaran_pabrik
      selisih     = disetor aktual - seharusnya_disetor
    """
    q = _range_query(date_from, date_to)
    txns = await db.transactions.find(q, {"_id": 0}).to_list(50000)
    exps = await db.expenses.find(q, {"_id": 0}).to_list(50000)
    pos = await db.purchase_orders.find(q, {"_id": 0}).to_list(50000)
    deps = await db.deposits.find(q, {"_id": 0}).to_list(50000)

    def blank(tgl):
        return {"tanggal": tgl, "penjualan": 0, "kas_masuk": 0, "kredit": 0, "free_hpp": 0,
                "hpp": 0, "laba_kotor": 0, "pembelian": 0, "pembayaran_pabrik": 0,
                "operasional": 0, "laba_bersih": 0, "seharusnya_disetor": 0,
                "disetor": 0, "selisih": 0, "jumlah_txn": 0}

    hari = {}
    for t in txns:
        d = hari.setdefault(t["tanggal"], blank(t["tanggal"]))
        total = float(t.get("total") or 0)
        hpp = txn_hpp(t)
        d["penjualan"] += total
        d["hpp"] += hpp
        d["jumlah_txn"] += 1
        if t.get("status") == "lunas":
            d["kas_masuk"] += total
        elif t.get("status") == "kredit":
            d["kredit"] += total
        elif t.get("status") == "free":
            d["free_hpp"] += hpp
    for e in exps:
        d = hari.setdefault(e["tanggal"], blank(e["tanggal"]))
        d["operasional"] += float(e.get("jumlah") or 0)
    for p in pos:
        d = hari.setdefault(p["tanggal"], blank(p["tanggal"]))
        tot = float(p.get("total") or 0)
        d["pembelian"] += tot
        if p.get("status") == "lunas":
            d["pembayaran_pabrik"] += tot
    for s in deps:
        d = hari.setdefault(s["tanggal"], blank(s["tanggal"]))
        d["disetor"] += float(s.get("jumlah") or 0)

    for d in hari.values():
        d["laba_kotor"] = d["penjualan"] - d["hpp"]
        d["laba_bersih"] = d["laba_kotor"] - d["operasional"]
        d["seharusnya_disetor"] = d["kas_masuk"] - d["operasional"] - d["pembayaran_pabrik"]
        d["selisih"] = d["disetor"] - d["seharusnya_disetor"]

    per_hari = [hari[k] for k in sorted(hari.keys())]
    keys = ["penjualan", "kas_masuk", "kredit", "free_hpp", "hpp", "laba_kotor", "pembelian",
            "pembayaran_pabrik", "operasional", "laba_bersih", "seharusnya_disetor",
            "disetor", "selisih", "jumlah_txn"]
    total = {k: sum(d[k] for d in per_hari) for k in keys}
    total["hari_aktif"] = len(per_hari)

    # Rincian modal barang free per alasan (promo / sampel / pemakaian sendiri / manual)
    per_alasan = {}
    for t in txns:
        if t.get("status") != "free":
            continue
        key = (t.get("free_alasan") or "").strip() or "Pemakaian Sendiri"
        row = per_alasan.setdefault(key, {"alasan": key, "hpp": 0, "jumlah_txn": 0, "qty": 0})
        row["hpp"] += txn_hpp(t)
        row["jumlah_txn"] += 1
        row["qty"] += sum(float(i.get("qty") or 0) for i in t.get("items", []))
    free_per_alasan = sorted(per_alasan.values(), key=lambda x: x["hpp"], reverse=True)

    return {"date_from": date_from, "date_to": date_to, "per_hari": per_hari, "total": total,
            "free_per_alasan": free_per_alasan}

def resolve_range(year, month, date_from, date_to):
    """Mode bulan (year+month) atau mode rentang tanggal bebas."""
    if date_from and date_to:
        if date_from > date_to:
            raise HTTPException(status_code=422, detail="Tanggal 'Dari' tidak boleh melebihi 'Sampai'")
        return date_from, date_to
    if year and month:
        if month < 1 or month > 12:
            raise HTTPException(status_code=422, detail="Bulan harus antara 1 dan 12")
        if year < 2000 or year > 2100:
            raise HTTPException(status_code=422, detail="Tahun tidak valid")
        last = calendar.monthrange(year, month)[1]
        return f"{year:04d}-{month:02d}-01", f"{year:04d}-{month:02d}-{last:02d}"
    raise HTTPException(status_code=422,
        detail="Pilih bulan (year & month) atau rentang tanggal (date_from & date_to)")

@api_router.get("/laporan/harian")
async def laporan_harian(year: Optional[int] = None, month: Optional[int] = None,
                         date_from: Optional[str] = None, date_to: Optional[str] = None,
                         user: dict = Depends(get_current_user)):
    df, dt = resolve_range(year, month, date_from, date_to)
    return await compute_harian(df, dt)

# ----------------------- Export -----------------------

def rupiah(n):
    val = int(round(float(n or 0)))
    prefix = "-Rp " if val < 0 else "Rp "
    return prefix + f"{abs(val):,}".replace(",", ".")

TXN_STATUS_LABEL = {"lunas": "Lunas", "kredit": "Kredit", "free": "Free"}

def status_label(s):
    return TXN_STATUS_LABEL.get(s, s or "-")

def status_text(t):
    """Label status + alasan bila barang free, mis. 'Free (Promo)'."""
    label = status_label(t.get("status"))
    alasan = (t.get("free_alasan") or "").strip()
    if t.get("status") == "free" and alasan:
        return f"{label} ({alasan})"
    return label

async def fetch_txns(date_from, date_to, customer_id, status):
    q = build_txn_query(date_from, date_to, customer_id, status)
    return await db.transactions.find(q, {"_id": 0}).sort("tanggal", 1).to_list(10000)

@api_router.get("/export/excel")
async def export_excel(date_from: Optional[str] = None, date_to: Optional[str] = None,
                       customer_id: Optional[str] = None, status: Optional[str] = None,
                       user: dict = Depends(get_current_user)):
    txns = await fetch_txns(date_from, date_to, customer_id, status)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rekap Penjualan"
    header_fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    headers = ["Tanggal", "Customer", "Produk", "Varian", "Harga SO", "Qty", "Subtotal",
               "HPP", "Laba", "Status"]
    ws.append(headers)
    for c in ws[1]:
        c.fill = header_fill
        c.font = header_font
    for t in txns:
        for it in t["items"]:
            sub = float(it.get("subtotal") or 0)
            hpp = float(it.get("hpp_subtotal") or 0)
            ws.append([t["tanggal"], t["customer_nama"], it["product_nama"], it["variant_label"],
                       it["harga_satuan"], it["qty"], sub, hpp, sub - hpp,
                       status_text(t)])
    total = sum(t["total"] for t in txns)
    total_hpp = sum(txn_hpp(t) for t in txns)
    ws.append([])
    ws.append(["", "", "", "", "", "", "TOTAL OMSET", total])
    ws.append(["", "", "", "", "", "", "TOTAL HPP", total_hpp])
    ws.append(["", "", "", "", "", "", "LABA KOTOR", total - total_hpp])
    for col in "ABCDEFGHIJ":
        ws.column_dimensions[col].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=rekap_penjualan.xlsx"})

@api_router.get("/export/pdf")
async def export_pdf(date_from: Optional[str] = None, date_to: Optional[str] = None,
                     customer_id: Optional[str] = None, status: Optional[str] = None,
                     user: dict = Depends(get_current_user)):
    txns = await fetch_txns(date_from, date_to, customer_id, status)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=18, bottomMargin=18,
                            leftMargin=18, rightMargin=18)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], fontSize=16, textColor=colors.HexColor("#0f172a"))
    elems = [Paragraph("Rekap Laporan Penjualan", title_style)]
    subtitle = f"Periode: {date_from or 'Semua'} s/d {date_to or 'Semua'}"
    elems.append(Paragraph(subtitle, styles["Normal"]))
    elems.append(Spacer(1, 10))
    data = [["Tanggal", "Customer", "Produk", "Varian", "Harga SO", "Qty", "Subtotal", "HPP", "Laba", "Status"]]
    for t in txns:
        for it in t["items"]:
            sub = float(it.get("subtotal") or 0)
            hpp = float(it.get("hpp_subtotal") or 0)
            data.append([t["tanggal"], t["customer_nama"], it["product_nama"], it["variant_label"],
                         rupiah(it["harga_satuan"]), str(int(float(it["qty"]))),
                         rupiah(sub), rupiah(hpp), rupiah(sub - hpp),
                         status_text(t)])
    total = sum(t["total"] for t in txns)
    total_hpp = sum(txn_hpp(t) for t in txns)
    data.append(["", "", "", "", "", "", rupiah(total), rupiah(total_hpp), rupiah(total - total_hpp), "TOTAL"])
    tbl = Table(data, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#f1f5f9")]),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#0f172a")),
        ("TEXTCOLOR", (0, -1), (-1, -1), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    elems.append(tbl)
    doc.build(elems)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=rekap_penjualan.pdf"})

BULAN_ID = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni",
            "Juli", "Agustus", "September", "Oktober", "November", "Desember"]

def _style_header(ws, ncols):
    header_fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for c in ws[1]:
        c.fill = header_fill
        c.font = header_font
    for i in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20

def _xlsx_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"})

def _pdf_table(data, col_widths=None):
    tbl = Table(data, repeatRows=1, colWidths=col_widths)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    return tbl

@api_router.get("/export/purchases/excel")
async def export_purchases_excel(date_from: Optional[str] = None, date_to: Optional[str] = None,
                                 status: Optional[str] = None, pabrik: Optional[str] = None,
                                 user: dict = Depends(get_current_user)):
    q = build_po_query(date_from, date_to, status, pabrik)
    pos = await db.purchase_orders.find(q, {"_id": 0}).sort("tanggal", 1).to_list(10000)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Laporan Pembelian"
    ws.append(["Tanggal", "No PO", "Pabrik", "Produk", "Varian", "Harga Pabrik", "Qty", "Subtotal", "Status Bayar"])
    _style_header(ws, 9)
    for p in pos:
        for it in p.get("items", []):
            ws.append([p["tanggal"], p.get("no_po", ""), p.get("pabrik", ""),
                       it.get("product_nama", ""), it.get("variant_label", ""),
                       float(it.get("harga_pabrik") or 0), float(it.get("qty") or 0),
                       float(it.get("subtotal") or 0),
                       "Sudah Bayar" if p.get("status") == "lunas" else "Belum Bayar"])
    total = sum(float(p.get("total") or 0) for p in pos)
    hutang = sum(float(p.get("total") or 0) for p in pos if p.get("status") == "belum_lunas")
    ws.append([])
    ws.append(["", "", "", "", "", "", "", "TOTAL PEMBELIAN", total])
    ws.append(["", "", "", "", "", "", "", "HUTANG KE PABRIK", hutang])
    return _xlsx_response(wb, "laporan_pembelian.xlsx")

@api_router.get("/export/purchases/pdf")
async def export_purchases_pdf(date_from: Optional[str] = None, date_to: Optional[str] = None,
                               status: Optional[str] = None, pabrik: Optional[str] = None,
                               user: dict = Depends(get_current_user)):
    q = build_po_query(date_from, date_to, status, pabrik)
    pos = await db.purchase_orders.find(q, {"_id": 0}).sort("tanggal", 1).to_list(10000)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=18, bottomMargin=18,
                            leftMargin=18, rightMargin=18)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], fontSize=16, textColor=colors.HexColor("#0f172a"))
    elems = [Paragraph("Laporan Pembelian ke Pabrik", title_style),
             Paragraph(f"Periode: {date_from or 'Semua'} s/d {date_to or 'Semua'}", styles["Normal"]),
             Spacer(1, 10)]
    data = [["Tanggal", "No PO", "Pabrik", "Produk", "Varian", "Harga", "Qty", "Subtotal", "Status"]]
    for p in pos:
        for it in p.get("items", []):
            data.append([p["tanggal"], p.get("no_po", "-"), p.get("pabrik", ""),
                         it.get("product_nama", ""), it.get("variant_label", ""),
                         rupiah(it.get("harga_pabrik") or 0), str(int(float(it.get("qty") or 0))),
                         rupiah(it.get("subtotal") or 0),
                         "Sudah Bayar" if p.get("status") == "lunas" else "Belum Bayar"])
    if len(data) == 1:
        data.append(["-"] * 9)
    elems.append(_pdf_table(data))
    total = sum(float(p.get("total") or 0) for p in pos)
    hutang = sum(float(p.get("total") or 0) for p in pos if p.get("status") == "belum_lunas")
    elems.append(Spacer(1, 12))
    elems.append(_pdf_table([["TOTAL PEMBELIAN", rupiah(total)], ["HUTANG KE PABRIK", rupiah(hutang)]]))
    doc.build(elems)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=laporan_pembelian.pdf"})

@api_router.get("/export/laba/excel")
async def export_laba_excel(year: int, month: int, user: dict = Depends(get_current_user)):
    if month < 1 or month > 12 or year < 2000 or year > 2100:
        raise HTTPException(status_code=422, detail="Periode tidak valid")
    r = await compute_laba(f"{year:04d}-{month:02d}")
    periode = f"{BULAN_ID[month]} {year}"
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Ringkasan Laba"
    ws.append(["Keterangan", "Jumlah"])
    _style_header(ws, 2)
    for label, val in [
        (f"Periode: {periode}", ""),
        ("Omset Penjualan", r["omset"]),
        ("HPP (Harga Pabrik Barang Terjual)", r["hpp"]),
        ("LABA KOTOR", r["laba_kotor"]),
        ("Pengeluaran Operasional", r["pengeluaran"]),
        ("LABA BERSIH", r["laba_bersih"]),
        ("Margin Kotor (%)", r["margin_kotor_pct"]),
        ("Margin Bersih (%)", r["margin_bersih_pct"]),
        ("Total PO ke Pabrik (bulan ini)", r["total_po_pabrik"]),
        ("Hutang ke Pabrik (belum dibayar)", r["hutang_pabrik"]),
    ]:
        ws.append([label, val])

    ws2 = wb.create_sheet("Laba per Produk")
    ws2.append(["Produk", "Qty Terjual", "Omset", "HPP", "Laba"])
    _style_header(ws2, 5)
    for p in r["per_product"]:
        ws2.append([p["nama"], p["qty"], p["omset"], p["hpp"], p["laba"]])

    ws3 = wb.create_sheet("Pengeluaran")
    ws3.append(["Kategori", "Jumlah"])
    _style_header(ws3, 2)
    for k in r["per_kategori"]:
        ws3.append([k["kategori"], k["jumlah"]])

    ws4 = wb.create_sheet("Harian")
    ws4.append(["Tanggal", "Omset", "HPP", "Pengeluaran", "Laba"])
    _style_header(ws4, 5)
    for d in r["per_day"]:
        ws4.append([d["tanggal"], d["omset"], d["hpp"], d["pengeluaran"], d["laba"]])

    return _xlsx_response(wb, f"laporan_laba_{year}_{month:02d}.xlsx")

@api_router.get("/export/laba/pdf")
async def export_laba_pdf(year: int, month: int, user: dict = Depends(get_current_user)):
    if month < 1 or month > 12 or year < 2000 or year > 2100:
        raise HTTPException(status_code=422, detail="Periode tidak valid")
    r = await compute_laba(f"{year:04d}-{month:02d}")
    periode = f"{BULAN_ID[month]} {year}"
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=20, bottomMargin=20,
                            leftMargin=24, rightMargin=24)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], fontSize=16, textColor=colors.HexColor("#0f172a"))
    h = ParagraphStyle("h", parent=styles["Heading3"], textColor=colors.HexColor("#1D4ED8"))
    elems = [Paragraph("Laporan Laba / Rugi", title_style),
             Paragraph(f"CV Citra Pangan Lestari — Periode {periode}", styles["Normal"]),
             Spacer(1, 12), Paragraph("Ringkasan", h)]
    elems.append(_pdf_table([
        ["Omset Penjualan", rupiah(r["omset"])],
        ["HPP (Harga Pabrik Barang Terjual)", "- " + rupiah(r["hpp"])],
        ["LABA KOTOR", rupiah(r["laba_kotor"])],
        ["Pengeluaran Operasional", "- " + rupiah(r["pengeluaran"])],
        ["LABA BERSIH", rupiah(r["laba_bersih"])],
        ["Margin Bersih", f'{r["margin_bersih_pct"]}%'],
        ["Total PO ke Pabrik (bulan ini)", rupiah(r["total_po_pabrik"])],
        ["Hutang ke Pabrik (belum dibayar)", rupiah(r["hutang_pabrik"])],
    ], col_widths=[280, 240]))

    elems.append(Spacer(1, 14))
    elems.append(Paragraph("Laba per Produk", h))
    data = [["Produk", "Qty", "Omset", "HPP", "Laba"]]
    for p in r["per_product"]:
        data.append([p["nama"], str(int(p["qty"])), rupiah(p["omset"]), rupiah(p["hpp"]), rupiah(p["laba"])])
    if len(data) == 1:
        data.append(["Belum ada penjualan", "-", "-", "-", "-"])
    elems.append(_pdf_table(data, col_widths=[180, 50, 100, 100, 100]))

    elems.append(Spacer(1, 14))
    elems.append(Paragraph("Pengeluaran Operasional", h))
    data2 = [["Kategori", "Jumlah"]]
    for k in r["per_kategori"]:
        data2.append([k["kategori"], rupiah(k["jumlah"])])
    if len(data2) == 1:
        data2.append(["Belum ada pengeluaran", "-"])
    elems.append(_pdf_table(data2, col_widths=[280, 240]))

    doc.build(elems)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=laporan_laba_{year}_{month:02d}.pdf"})

@api_router.get("/export/laba/yearly/excel")
async def export_laba_yearly_excel(year: int, user: dict = Depends(get_current_user)):
    if year < 2000 or year > 2100:
        raise HTTPException(status_code=422, detail="Tahun tidak valid")
    r = await compute_laba_tahunan(year)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Laba {year}"
    ws.append(["Bulan", "Omset", "HPP", "Laba Kotor", "Pengeluaran", "Laba Bersih",
               "Margin Bersih (%)", "Transaksi", "PO Pabrik"])
    _style_header(ws, 9)
    for b in r["per_bulan"]:
        ws.append([b["nama_bulan"], b["omset"], b["hpp"], b["laba_kotor"], b["pengeluaran"],
                   b["laba_bersih"], b["margin_bersih_pct"], b["total_txn"], b["po_pabrik"]])
    ws.append([])
    ws.append(["TOTAL SETAHUN", r["total_omset"], r["total_hpp"], r["total_laba_kotor"],
               r["total_pengeluaran"], r["total_laba_bersih"], r["margin_bersih_pct"],
               r["total_txn"], r["total_po_pabrik"]])
    ws.append([])
    if r["bulan_terbaik"]:
        ws.append(["Bulan terbaik", r["bulan_terbaik"]["nama_bulan"], r["bulan_terbaik"]["laba_bersih"]])
        ws.append(["Bulan terburuk", r["bulan_terburuk"]["nama_bulan"], r["bulan_terburuk"]["laba_bersih"]])
        ws.append(["Rata-rata laba bersih per bulan aktif", r["rata_laba_bersih"]])
    return _xlsx_response(wb, f"laporan_tahunan_{year}.xlsx")

@api_router.get("/export/laba/yearly/pdf")
async def export_laba_yearly_pdf(year: int, user: dict = Depends(get_current_user)):
    if year < 2000 or year > 2100:
        raise HTTPException(status_code=422, detail="Tahun tidak valid")
    r = await compute_laba_tahunan(year)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=20, bottomMargin=20,
                            leftMargin=24, rightMargin=24)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], fontSize=16, textColor=colors.HexColor("#0f172a"))
    h = ParagraphStyle("h", parent=styles["Heading3"], textColor=colors.HexColor("#1D4ED8"))
    elems = [Paragraph(f"Laporan Laba Tahunan {year}", title_style),
             Paragraph("CV Citra Pangan Lestari — perbandingan laba antar bulan", styles["Normal"]),
             Spacer(1, 12)]
    data = [["Bulan", "Omset", "HPP", "Laba Kotor", "Pengeluaran", "Laba Bersih", "Margin"]]
    for b in r["per_bulan"]:
        data.append([b["nama_bulan"], rupiah(b["omset"]), rupiah(b["hpp"]), rupiah(b["laba_kotor"]),
                     rupiah(b["pengeluaran"]), rupiah(b["laba_bersih"]), f'{b["margin_bersih_pct"]}%'])
    data.append(["TOTAL", rupiah(r["total_omset"]), rupiah(r["total_hpp"]), rupiah(r["total_laba_kotor"]),
                 rupiah(r["total_pengeluaran"]), rupiah(r["total_laba_bersih"]),
                 f'{r["margin_bersih_pct"]}%'])
    elems.append(_pdf_table(data, col_widths=[80, 105, 105, 105, 105, 105, 60]))

    if r["bulan_terbaik"]:
        elems.append(Spacer(1, 14))
        elems.append(Paragraph("Catatan", h))
        elems.append(_pdf_table([
            ["Bulan paling untung", f'{r["bulan_terbaik"]["nama_bulan"]} — {rupiah(r["bulan_terbaik"]["laba_bersih"])}'],
            ["Bulan paling rendah", f'{r["bulan_terburuk"]["nama_bulan"]} — {rupiah(r["bulan_terburuk"]["laba_bersih"])}'],
            ["Rata-rata laba bersih / bulan aktif", rupiah(r["rata_laba_bersih"])],
        ], col_widths=[260, 300]))

    doc.build(elems)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=laporan_tahunan_{year}.pdf"})

@api_router.get("/export/harian/excel")
async def export_harian_excel(year: Optional[int] = None, month: Optional[int] = None,
                              date_from: Optional[str] = None, date_to: Optional[str] = None,
                              user: dict = Depends(get_current_user)):
    df, dt = resolve_range(year, month, date_from, date_to)
    r = await compute_harian(df, dt)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Laporan Harian"
    ws.append(["Tanggal", "Penjualan", "Pembelian", "Operasional", "Laba Bersih",
               "Seharusnya Disetor", "Disetor", "Selisih"])
    _style_header(ws, 8)
    for d in r["per_hari"]:
        ws.append([d["tanggal"], d["penjualan"], d["pembelian"], d["operasional"],
                   d["laba_bersih"], d["seharusnya_disetor"], d["disetor"], d["selisih"]])
    t = r["total"]
    ws.append([])
    ws.append(["TOTAL", t["penjualan"], t["pembelian"], t["operasional"], t["laba_bersih"],
               t["seharusnya_disetor"], t["disetor"], t["selisih"]])

    ws2 = wb.create_sheet("Rincian")
    ws2.append(["Tanggal", "Jml Transaksi", "Penjualan", "Kas Masuk (Lunas)", "Kredit",
                "HPP", "Laba Kotor", "Modal Barang Free", "Pembelian PO",
                "Pembayaran ke Pabrik", "Operasional", "Laba Bersih"])
    _style_header(ws2, 12)
    for d in r["per_hari"]:
        ws2.append([d["tanggal"], d["jumlah_txn"], d["penjualan"], d["kas_masuk"], d["kredit"],
                    d["hpp"], d["laba_kotor"], d["free_hpp"], d["pembelian"],
                    d["pembayaran_pabrik"], d["operasional"], d["laba_bersih"]])
    ws2.append([])
    ws2.append(["TOTAL", t["jumlah_txn"], t["penjualan"], t["kas_masuk"], t["kredit"],
                t["hpp"], t["laba_kotor"], t["free_hpp"], t["pembelian"],
                t["pembayaran_pabrik"], t["operasional"], t["laba_bersih"]])

    if r.get("free_per_alasan"):
        ws3 = wb.create_sheet("Barang Free")
        ws3.append(["Alasan", "Jumlah Transaksi", "Total Qty", "Modal (HPP)"])
        _style_header(ws3, 4)
        for f in r["free_per_alasan"]:
            ws3.append([f["alasan"], f["jumlah_txn"], f["qty"], f["hpp"]])
        ws3.append([])
        ws3.append(["TOTAL", sum(f["jumlah_txn"] for f in r["free_per_alasan"]),
                    sum(f["qty"] for f in r["free_per_alasan"]), t["free_hpp"]])

    return _xlsx_response(wb, f"laporan_harian_{df}_sd_{dt}.xlsx")

@api_router.get("/export/harian/pdf")
async def export_harian_pdf(year: Optional[int] = None, month: Optional[int] = None,
                            date_from: Optional[str] = None, date_to: Optional[str] = None,
                            user: dict = Depends(get_current_user)):
    df, dt = resolve_range(year, month, date_from, date_to)
    r = await compute_harian(df, dt)
    t = r["total"]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=20, bottomMargin=20,
                            leftMargin=24, rightMargin=24)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], fontSize=16, textColor=colors.HexColor("#0f172a"))
    h = ParagraphStyle("h", parent=styles["Heading3"], textColor=colors.HexColor("#1D4ED8"))
    small = ParagraphStyle("s", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#475569"))
    elems = [Paragraph("Laporan Harian", title_style),
             Paragraph(f"CV Citra Pangan Lestari — {df} s/d {dt}", styles["Normal"]),
             Spacer(1, 12)]

    data = [["Tanggal", "Penjualan", "Pembelian", "Operasional", "Laba Bersih",
             "Seharusnya Disetor", "Disetor", "Selisih"]]
    for d in r["per_hari"]:
        data.append([d["tanggal"], rupiah(d["penjualan"]), rupiah(d["pembelian"]),
                     rupiah(d["operasional"]), rupiah(d["laba_bersih"]),
                     rupiah(d["seharusnya_disetor"]), rupiah(d["disetor"]), rupiah(d["selisih"])])
    if len(data) == 1:
        data.append(["Belum ada data", "-", "-", "-", "-", "-", "-", "-"])
    data.append(["TOTAL", rupiah(t["penjualan"]), rupiah(t["pembelian"]), rupiah(t["operasional"]),
                 rupiah(t["laba_bersih"]), rupiah(t["seharusnya_disetor"]),
                 rupiah(t["disetor"]), rupiah(t["selisih"])])
    elems.append(_pdf_table(data, col_widths=[70, 95, 95, 90, 95, 105, 95, 95]))

    elems.append(Spacer(1, 14))
    elems.append(Paragraph("Ringkasan Periode", h))
    elems.append(_pdf_table([
        ["Total Penjualan (Lunas + Kredit)", rupiah(t["penjualan"])],
        ["Kas Masuk (hanya Lunas)", rupiah(t["kas_masuk"])],
        ["Penjualan Kredit (belum dibayar)", rupiah(t["kredit"])],
        ["Modal Barang Free / Pemakaian Sendiri", rupiah(t["free_hpp"])],
        ["HPP", rupiah(t["hpp"])],
        ["Laba Kotor", rupiah(t["laba_kotor"])],
        ["Pengeluaran Operasional", rupiah(t["operasional"])],
        ["LABA BERSIH", rupiah(t["laba_bersih"])],
        ["Pembelian ke Pabrik", rupiah(t["pembelian"])],
        ["Pembayaran ke Pabrik", rupiah(t["pembayaran_pabrik"])],
        ["Seharusnya Disetor", rupiah(t["seharusnya_disetor"])],
        ["Sudah Disetor", rupiah(t["disetor"])],
        ["SELISIH", rupiah(t["selisih"])],
    ], col_widths=[300, 260]))

    if r.get("free_per_alasan"):
        elems.append(Spacer(1, 14))
        elems.append(Paragraph("Rincian Barang Free", h))
        fdata = [["Alasan", "Jml Transaksi", "Total Qty", "Modal (HPP)"]]
        for f in r["free_per_alasan"]:
            fdata.append([f["alasan"], str(f["jumlah_txn"]), str(int(f["qty"])), rupiah(f["hpp"])])
        fdata.append(["TOTAL", str(sum(f["jumlah_txn"] for f in r["free_per_alasan"])),
                      str(int(sum(f["qty"] for f in r["free_per_alasan"]))), rupiah(t["free_hpp"])])
        elems.append(_pdf_table(fdata, col_widths=[240, 110, 110, 140]))

    elems.append(Spacer(1, 10))
    elems.append(Paragraph(
        "Catatan: Seharusnya Disetor = Kas Masuk (penjualan Lunas) − Operasional − Pembayaran ke Pabrik. "
        "Penjualan Kredit belum dihitung sebagai kas masuk. Barang Free tidak menghasilkan omset, "
        "tetapi modal pabriknya tetap dihitung sebagai biaya.", small))

    doc.build(elems)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=laporan_harian_{df}_sd_{dt}.pdf"})

# ----------------------- Seed -----------------------

SEED_PRODUCTS = [
    ("CHEERS ALKALINE", ["1200", "550", "330", "230", "GALON"]),
    ("CHEERS REGULAR", ["1500", "600", "330", "220", "GALON"]),
    ("VEMA", ["1500", "600", "220", "GALON"]),
]
DEFAULT_PRICES = {
    "1200": (45000, 47000, 50000), "1500": (40000, 42000, 45000),
    "600": (30000, 32000, 35000), "550": (30000, 32000, 35000),
    "330": (25000, 27000, 30000), "230": (20000, 22000, 25000),
    "220": (20000, 22000, 24000), "GALON": (18000, 19000, 20000),
}

async def migrate_harga_pabrik():
    """Tambah field harga_pabrik pada varian produk lama (default 0).

    Nilainya sengaja 0, bukan angka karangan. Halaman Laba akan menampilkan
    peringatan sampai pemilik mengisi harga pabrik yang sebenarnya di menu Produk.
    """
    updated = 0
    async for p in db.products.find({}):
        variants = p.get("variants", [])
        changed = False
        for v in variants:
            if "harga_pabrik" not in v:
                v["harga_pabrik"] = 0
                changed = True
        if changed:
            await db.products.update_one({"_id": p["_id"]}, {"$set": {"variants": variants}})
            updated += 1
    if updated:
        logger.info(f"Migrasi harga_pabrik: {updated} produk diperbarui")

async def migrate_transaction_hpp():
    """Lengkapi transaksi lama dengan snapshot HPP memakai harga pabrik saat ini."""
    prods = await db.products.find({}, {"_id": 0}).to_list(1000)
    pmap = {p["id"]: p for p in prods}
    updated = 0
    async for t in db.transactions.find({"total_hpp": {"$exists": False}}):
        items = t.get("items", [])
        for it in items:
            hp = float(it.get("harga_pabrik") or 0)
            if hp <= 0:
                p = pmap.get(it.get("product_id"))
                if p:
                    v = next((x for x in p.get("variants", []) if x.get("label") == it.get("variant_label")), None)
                    if v:
                        hp = float(v.get("harga_pabrik") or 0)
            it["harga_pabrik"] = hp
            it["hpp_subtotal"] = hp * float(it.get("qty") or 0)
        total_hpp = sum(float(i.get("hpp_subtotal") or 0) for i in items)
        await db.transactions.update_one({"_id": t["_id"]}, {"$set": {
            "items": items, "total_hpp": total_hpp,
            "laba_kotor": float(t.get("total") or 0) - total_hpp}})
        updated += 1
    if updated:
        logger.info(f"Migrasi HPP transaksi: {updated} transaksi diperbarui")

async def migrate_single_price():
    """Sederhanakan ke satu harga jual (Harga SO).

    Menghapus harga_grosir & harga_retail dari varian produk, default_price_type
    dari customer, dan price_type dari item transaksi. Harga SO dipertahankan.
    """
    prod = await db.products.update_many(
        {}, {"$unset": {"variants.$[].harga_grosir": "", "variants.$[].harga_retail": ""}})
    cust = await db.customers.update_many({}, {"$unset": {"default_price_type": ""}})
    txn = await db.transactions.update_many({}, {"$unset": {"items.$[].price_type": ""}})
    if prod.modified_count or cust.modified_count or txn.modified_count:
        logger.info(f"Migrasi harga tunggal: {prod.modified_count} produk, "
                    f"{cust.modified_count} customer, {txn.modified_count} transaksi")

async def migrate_status_kredit():
    """Ganti status transaksi 'belum_lunas' menjadi 'kredit' (istilah baru)."""
    res = await db.transactions.update_many({"status": "belum_lunas"}, {"$set": {"status": "kredit"}})
    if res.modified_count:
        logger.info(f"Migrasi status transaksi ke 'kredit': {res.modified_count} transaksi")

@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.login_attempts.create_index("identifier", unique=True)
    await db.purchase_orders.create_index("tanggal")
    await db.purchase_orders.create_index("status")
    await db.expenses.create_index("tanggal")
    await db.expenses.create_index("kategori")
    await db.deposits.create_index("tanggal")
    admin_email = os.environ.get("ADMIN_EMAIL", "admin@example.com").lower()
    admin_password = os.environ.get("ADMIN_PASSWORD", "admin123")
    # Seed admin only when no user exists yet. Do NOT overwrite an existing
    # user's email/password on restart so self-service profile edits persist.
    if await db.users.count_documents({}) == 0:
        await db.users.insert_one({"email": admin_email, "password_hash": hash_password(admin_password),
                                   "name": "Admin", "role": "admin",
                                   "created_at": datetime.now(timezone.utc).isoformat()})
        logger.info("Admin dibuat")

    if await db.products.count_documents({}) == 0 and await db.meta.count_documents({"key": "products_seeded"}) == 0:
        for nama, labels in SEED_PRODUCTS:
            variants = []
            for lb in labels:
                g, s, r = DEFAULT_PRICES.get(lb, (0, 0, 0))
                variants.append({"label": lb, "harga_pabrik": 0, "harga_so": s})
            await db.products.insert_one({"id": str(uuid.uuid4()), "nama": nama, "variants": variants,
                                          "created_at": datetime.now(timezone.utc).isoformat()})
        await db.meta.insert_one({"key": "products_seeded",
                                  "at": datetime.now(timezone.utc).isoformat()})
        logger.info("Produk awal dibuat")

    await migrate_harga_pabrik()
    await migrate_transaction_hpp()
    await migrate_single_price()
    await migrate_status_kredit()

app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=[os.environ.get("FRONTEND_URL", "http://localhost:3000"), "http://localhost:3000"],
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
